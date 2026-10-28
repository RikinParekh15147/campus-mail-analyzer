"""Bounded attachment text extraction with source-aligned layout locators."""

from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
import hashlib


class TextBuilder:
    def __init__(self):
        self.parts = []
        self.length = 0
        self.locators = []

    def append(self, text, **location):
        start = self.length
        self.parts.append(text)
        self.length += len(text)
        if location and text:
            self.locators.append({**location, "start": start, "end": self.length})

    def text(self):
        return "".join(self.parts)


class VisibleText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.builder = TextBuilder()
        self.parts = self.builder.parts
        self.skip = 0
        self.table_index = 0
        self.row_index = 0
        self.row_start = None
        self.cell_index = 0

    @property
    def locators(self):
        return self.builder.locators

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1
            return
        if self.skip:
            return
        if tag in ("br", "p", "div", "li"):
            self.builder.append("\n")
        if tag == "table":
            self.table_index += 1
            self.row_index = 0
            self.builder.append("\n")
        if tag == "tr":
            self.row_index += 1
            self.cell_index = 0
            self.builder.append("\n")
            self.row_start = self.builder.length
        if tag in ("td", "th"):
            if self.cell_index:
                self.builder.append("\t")
            self.cell_index += 1
        if tag == "a":
            url = dict(attrs).get("href", "")
            if url.startswith(("https://", "http://")):
                self.builder.append(" " + url + " ")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip = max(0, self.skip - 1)
        elif tag == "tr" and self.row_start is not None and not self.skip:
            if self.builder.length > self.row_start:
                self.locators.append({"kind": "html_row", "table": self.table_index,
                                      "row": self.row_index, "start": self.row_start,
                                      "end": self.builder.length})
            self.row_start = None

    def handle_data(self, data):
        if not self.skip:
            self.builder.append(data)


def text_quality(text):
    replacements = text.count("\ufffd")
    controls = sum(1 for char in text if ord(char) < 32 and char not in "\n\r\t")
    ratio = (replacements + controls) / max(len(text), 1)
    return {"replacement_chars": replacements, "control_chars": controls,
            "suspect_ratio": round(ratio, 6), "characters": len(text),
            "status": "needs_text_review" if ratio >= 0.01 else "usable"}


def decode_bytes(data, charset=None):
    encoding = charset or "utf-8-sig"
    try:
        return data.decode(encoding), None
    except LookupError:
        return data.decode("utf-8-sig", errors="replace"), "unsupported_charset"
    except UnicodeDecodeError:
        return data.decode(encoding, errors="replace"), "charset_decode_replacement"


def read_document(data, filename, charset=None):
    if len(data) > 20 * 1024 * 1024:
        raise ValueError("Attachment exceeds 20 MB limit")
    result = {
        "filename": Path(filename).name,
        "sha256": hashlib.sha256(data).hexdigest(),
        "text": "",
        "status": "unsupported",
        "locators": [],
        "text_quality": None,
        "decoding_warnings": [],
    }
    ext = Path(filename).suffix.lower()
    try:
        if ext in (".txt", ".csv"):
            result["text"], warning = decode_bytes(data, charset)
            if warning:
                result["decoding_warnings"].append(warning)
        elif ext in (".html", ".htm"):
            parser = VisibleText()
            decoded, warning = decode_bytes(data, charset)
            if warning:
                result["decoding_warnings"].append(warning)
            parser.feed(decoded)
            result["text"] = parser.builder.text()
            result["locators"] = parser.locators
        elif ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("Password protected")
            if len(reader.pages) > 100:
                raise ValueError("PDF exceeds 100-page processing limit")
            builder = TextBuilder()
            for index, page in enumerate(reader.pages, 1):
                if index > 1:
                    builder.append("\n")
                builder.append(page.extract_text() or "", kind="pdf_page", page=index)
            result["text"] = builder.text()
            result["locators"] = builder.locators
        elif ext == ".docx":
            from docx import Document
            from docx.table import Table

            document = Document(BytesIO(data))
            builder = TextBuilder()
            block_index = 0
            for block in document.iter_inner_content():
                if isinstance(block, Table):
                    for row_index, row in enumerate(block.rows, 1):
                        if builder.length:
                            builder.append("\n")
                        block_index += 1
                        builder.append("\t".join(cell.text for cell in row.cells),
                                       kind="docx_table_row", block=block_index,
                                       row=row_index)
                else:
                    if builder.length:
                        builder.append("\n")
                    block_index += 1
                    builder.append(block.text, kind="docx_paragraph", block=block_index)
            result["text"] = builder.text()
            result["locators"] = builder.locators
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter

            book = load_workbook(BytesIO(data), read_only=True, data_only=True)
            try:
                builder = TextBuilder()
                for sheet in book:
                    if sheet.max_row > 10000:
                        raise ValueError("Spreadsheet exceeds 10,000 rows")
                    if builder.length:
                        builder.append("\n")
                    builder.append("[SHEET: " + sheet.title + "]\n")
                    for row_index, row in enumerate(sheet.values, 1):
                        if row_index > 1:
                            builder.append("\n")
                        row_start = builder.length
                        for column_index, cell in enumerate(row, 1):
                            if column_index > 1:
                                builder.append("\t")
                            value = "" if cell is None else str(cell)
                            builder.append(value, kind="xlsx_cell", sheet=sheet.title,
                                           row=row_index,
                                           cell=f"{get_column_letter(column_index)}{row_index}")
                        if builder.length > row_start:
                            builder.locators.append({"kind": "xlsx_row",
                                                     "sheet": sheet.title,
                                                     "row": row_index,
                                                     "start": row_start,
                                                     "end": builder.length})
                result["text"] = builder.text()
                result["locators"] = builder.locators
            finally:
                book.close()
        else:
            return result
        quality = text_quality(result["text"])
        result["text_quality"] = quality
        result["status"] = (
            "needs_ocr_or_review" if not result["text"].strip()
            else "needs_text_review" if result["decoding_warnings"]
            or quality["status"] != "usable"
            else "readable"
        )
    except Exception as error:
        result.update(status="unreadable", error=str(error)[:300])
    return result
