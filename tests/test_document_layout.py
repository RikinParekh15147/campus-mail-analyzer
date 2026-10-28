"""PIPE-03 layout offsets, charset and poor-text review boundaries."""

import base64
from io import BytesIO

from docx import Document
from openpyxl import Workbook

from campusmail.documents import read_document
from campusmail.engine import Engine
from campusmail.ingestion import parse_message
from campusmail.store import Store


def test_html_table_rows_keep_separate_cells_and_source_offsets():
    data = (b"<table><tr><th>Course</th><th>Deadline</th></tr>"
            b"<tr><td>NLP201</td><td>25 October 2026</td></tr></table>")
    result = read_document(data, "notice.html")
    assert result["status"] == "readable"
    rows = [item for item in result["locators"] if item["kind"] == "html_row"]
    assert len(rows) == 2
    assert result["text"][rows[1]["start"]:rows[1]["end"]] == \
        "NLP201\t25 October 2026"


def test_docx_body_order_and_spreadsheet_cell_locators():
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "NLP201"
    table.cell(0, 1).text = "25 October 2026"
    document.add_paragraph("After table")
    stream = BytesIO()
    document.save(stream)
    result = read_document(stream.getvalue(), "notice.docx")
    assert result["status"] == "readable"
    assert result["text"] == "Before table\nNLP201\t25 October 2026\nAfter table"
    assert [item["kind"] for item in result["locators"]] == [
        "docx_paragraph", "docx_table_row", "docx_paragraph"]

    book = Workbook()
    sheet = book.active
    sheet.title = "Schedule"
    sheet.append(["Course", "Deadline"])
    sheet.append(["NLP201", "25 October 2026"])
    stream = BytesIO()
    book.save(stream)
    result = read_document(stream.getvalue(), "schedule.xlsx")
    assert result["status"] == "readable"
    cell = next(item for item in result["locators"]
                if item.get("cell") == "B2")
    assert result["text"][cell["start"]:cell["end"]] == "25 October 2026"
    assert any(item["kind"] == "xlsx_row" and item["row"] == 2
               for item in result["locators"])


def test_attachment_date_assertion_keeps_spreadsheet_cell_locator(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Tasks"
    sheet["B2"] = "NLP201 Assignment 2 must submit by 25 October 2026."
    stream = BytesIO()
    book.save(stream)
    store = Store(tmp_path / "db")
    owner = store.create_user("Student")["id"]
    resource = Engine(store).attach(owner, stream.getvalue(), "tasks.xlsx")
    assert resource["status"] == "readable"
    assignment = store.list(owner, "assignment")[0]
    evidence = assignment["field_evidence"]["due_date"]["evidence"]
    assert evidence["locator"]["kind"] == "xlsx_cell"
    assert evidence["locator"]["cell"] == "B2"


def test_pdf_page_offsets_are_bound_to_combined_text(monkeypatch):
    import pypdf

    class Page:
        def __init__(self, text):
            self.text = text

        def extract_text(self):
            return self.text

    class Reader:
        is_encrypted = False
        pages = [Page("First page"), Page("Second page")]

    monkeypatch.setattr(pypdf, "PdfReader", lambda stream: Reader())
    result = read_document(b"fixture", "notice.pdf")
    assert result["text"] == "First page\nSecond page"
    assert [(item["page"], result["text"][item["start"]:item["end"]])
            for item in result["locators"]] == [
                (1, "First page"), (2, "Second page")]


def test_declared_charset_decodes_and_corrupt_text_requires_review(tmp_path):
    encoded = base64.urlsafe_b64encode("Café registration".encode("iso-8859-1")).decode()
    raw = {"id": "mail-1", "payload": {"mimeType": "text/plain",
           "headers": [{"name": "Content-Type",
                        "value": "text/plain; charset=iso-8859-1"}],
           "body": {"data": encoded}}}
    parsed = parse_message(raw)
    assert parsed["text"] == "Café registration"
    assert parsed["text_quality_status"] == "usable"
    raw["payload"]["parts"] = [{"filename": "notice.txt", "mimeType": "text/plain",
                                   "headers": [{"name": "Content-Type",
                                                "value": "text/plain; charset=iso-8859-1"}],
                                   "body": {"attachmentId": "attachment-1"}}]
    assert parse_message(raw)["attachments"][0]["charset"] == "iso-8859-1"
    assert read_document("Café registration".encode("iso-8859-1"),
                         "notice.txt", "iso-8859-1")["status"] == "readable"

    corrupt = read_document(b"Due \xff\xfe tomorrow", "notice.txt")
    assert corrupt["status"] == "needs_text_review"
    assert corrupt["text_quality"]["replacement_chars"] == 2
    store = Store(tmp_path / "db")
    owner = store.create_user("Student")["id"]
    resource = Engine(store).attach(owner, b"NLP201 Assignment 2 \xff\xfe due tomorrow",
                                    "bad.txt")
    assert resource["status"] == "needs_text_review"
    assert "content_analysis" not in resource
    assert store.list(owner, "assignment") == []
    message = store.put(owner, "message", {
        "subject": "NLP201 Assignment 2", "text": corrupt["text"],
        "text_quality_status": corrupt["status"],
    }, "poor-text")
    result = Engine(store).process(owner, message)
    assert result["category"] == "unknown"
    assert store.list(owner, "assignment") == []
    assert any(item["source_key"].startswith("text-quality:")
               for item in store.list(owner, "review"))
